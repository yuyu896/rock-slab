"""
Tests for inventory-checklist-safety：空清单开始拦截、模板填写指引、导入空清单根因提示。
"""
import io

import openpyxl
import pytest
from rest_framework import status

from apps.assets.models import FixedAsset
from apps.assets.services import ledger
from apps.categories.models import Category
from apps.inventories.models import InventoryTask, InventoryInstanceItem, InventoryItem
from apps.organizations.models import Branch


def _item(code, management_type='instance'):
    item, _ = Category.objects.get_or_create(
        asset_code=code,
        defaults={
            'asset_category': '电子设备', 'item_category': '电脑',
            'asset_name': f'品目 {code}', 'unit': '台',
            'management_type': management_type,
        },
    )
    return item


def _seed_in_use(branch, item, n):
    for i in range(n):
        FixedAsset.objects.create(
            item=item, 内部编号=f'{item.asset_code}-{branch.code}-{i + 1}',
            当前状态='在用', branch=branch, 使用人='张三',
        )


def _task(branch, kind='instance', category=None):
    return InventoryTask.objects.create(
        name='安全测试盘', kind=kind, branch=branch, category=category,
        stock_bin='stock', missed_rule='zero', repeat_rule='last',
    )


def _make_xlsx(headers, rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in (rows or []):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = 'test.xlsx'
    return buf


@pytest.mark.django_db
class TestEmptyScopeStartGuard:
    def test_instance_scope_empty_start_rejected(self, authenticated_client, branch):
        _item('SG-1')  # 有品目但无在用实例
        task = _task(branch, kind='instance')
        resp = authenticated_client.post(f'/api/inventories/{task.id}/start')
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '无可盘对象' in str(resp.data['detail'])
        assert '实例档案' in str(resp.data['detail'])
        task.refresh_from_db()
        assert task.status == 'pending'
        assert task.instance_items.count() == 0

    def test_ledger_scope_empty_start_rejected(self, authenticated_client, branch):
        task = _task(branch, kind='stock')
        resp = authenticated_client.post(f'/api/inventories/{task.id}/start')
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '无可盘对象' in str(resp.data['detail'])
        task.refresh_from_db()
        assert task.status == 'pending'

    def test_scope_nonempty_start_ok(self, authenticated_client, branch):
        item = _item('SG-2')
        _seed_in_use(branch, item, 3)
        task = _task(branch, kind='instance')
        resp = authenticated_client.post(f'/api/inventories/{task.id}/start')
        assert resp.status_code == status.HTTP_200_OK
        task.refresh_from_db()
        assert task.status == 'in_progress'
        assert task.instance_items.count() == 3

    def test_category_filter_narrows_to_empty(self, authenticated_client, branch):
        """类目过滤后为空同样被拦（即便分公司有其他类目在用实例）。"""
        item_a = _item('SG-3A')
        item_b = _item('SG-3B')
        Category.objects.filter(pk=item_b.pk).update(asset_category='办公家具')
        _seed_in_use(branch, item_a, 2)
        task = _task(branch, kind='instance', category=item_b)
        resp = authenticated_client.post(f'/api/inventories/{task.id}/start')
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '类目' in str(resp.data['detail'])


@pytest.mark.django_db
class TestTemplateGuidance:
    def _start_task(self, client, branch):
        item = _item('SG-4')
        _seed_in_use(branch, item, 2)
        task = _task(branch, kind='instance')
        assert client.post(f'/api/inventories/{task.id}/start').status_code == 200
        return task

    def test_instance_template_has_dropdown_and_comments(self, authenticated_client, branch):
        task = self._start_task(authenticated_client, branch)
        resp = authenticated_client.get(f'/api/inventories/{task.id}/import-template')
        assert resp.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        dvs = list(ws.data_validations.dataValidation)
        assert len(dvs) == 1 and dvs[0].formula1 == '"已找到,未找到"'
        assert 'I2:I3' in str(dvs[0].sqref)
        assert ws['I1'].comment is not None and '已找到' in ws['I1'].comment.text
        assert ws['J1'].comment is not None
        assert ws.max_row == 3  # 表头 + 2 台，无示例行
        assert ws['F2'].value in ('在库', '在用')  # 状态列有值

    def test_ledger_template_has_qty_comment(self, authenticated_client, branch):
        item = _item('SG-5', 'quantity')
        ledger.apply_adjustment(branch, item, ledger.COLUMN_STOCK, 4, '测试造数')
        task = _task(branch, kind='stock')
        assert authenticated_client.post(f'/api/inventories/{task.id}/start').status_code == 200
        resp = authenticated_client.get(f'/api/inventories/{task.id}/import-template')
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        assert ws['F1'].comment is not None and '整数' in ws['F1'].comment.text


@pytest.mark.django_db
class TestImportHint:
    NEW_HEADERS = ['序号', '内部编号', '序列号', '品目编号', '品目名称', '状态', '使用人', '所属部门', '核对结果', '备注']

    def _upload(self, client, task):
        buf = _make_xlsx(
            self.NEW_HEADERS,
            [[1, 'X-1', '', 'X', '品目', '在用', '张三', '', '已找到', '']],
        )
        return client.post(
            f'/api/inventories/{task.id}/import-result', {'file': buf}, format='multipart',
        )

    def test_empty_checklist_import_returns_hint(self, authenticated_client, branch):
        item = _item('SG-6')
        inst = FixedAsset.objects.create(
            item=item, 内部编号='X-1', 当前状态='在用', branch=branch)
        task = _task(branch, kind='instance')
        InventoryInstanceItem.objects.create(task=task, instance=inst)
        task.status = 'in_progress'
        task.save(update_fields=['status'])
        InventoryInstanceItem.objects.filter(task=task).delete()  # 模拟历史空清单任务
        resp = self._upload(authenticated_client, task)
        assert resp.status_code == 200
        assert resp.data['imported'] == 0
        assert '清单为空' in resp.data['hint']
        assert '作废重建' in resp.data['hint']

    def test_nonempty_checklist_no_hint(self, authenticated_client, branch):
        item = _item('SG-7')
        inst = FixedAsset.objects.create(
            item=item, 内部编号='X-1', 当前状态='在用', branch=branch)
        task = _task(branch, kind='instance')
        InventoryInstanceItem.objects.create(task=task, instance=inst)
        task.status = 'in_progress'
        task.save(update_fields=['status'])
        resp = self._upload(authenticated_client, task)
        assert resp.status_code == 200
        assert resp.data['imported'] == 1
        assert 'hint' not in resp.data

    def test_pasted_invalid_verdict_still_rejected(self, authenticated_client, branch):
        item = _item('SG-8')
        inst = FixedAsset.objects.create(
            item=item, 内部编号='X-1', 当前状态='在用', branch=branch)
        task = _task(branch, kind='instance')
        InventoryInstanceItem.objects.create(task=task, instance=inst)
        task.status = 'in_progress'
        task.save(update_fields=['status'])
        buf = _make_xlsx(
            self.NEW_HEADERS,
            [[1, 'X-1', '', 'X', '品目', '在用', '张三', '', '没找到', '']],
        )
        resp = authenticated_client.post(
            f'/api/inventories/{task.id}/import-result', {'file': buf}, format='multipart')
        assert resp.status_code == 200
        assert any('已找到/未找到' in e for e in resp.data['errors'])
