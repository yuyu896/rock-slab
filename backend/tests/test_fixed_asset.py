"""实例档案（P2 第二刀）：冻结只读 + 序列号补录 + 生平查询 + 待补录筛选。

对应 fixed-asset-instance / document-instance-binding 能力。
"""
import datetime

import pytest
from conftest import _client_for

from apps.assets.views import FixedAssetViewSet

FA_EXPORT_HEADERS = FixedAssetViewSet.FA_EXPORT_HEADERS


def _grant(user, code):
    from apps.permissions.models import OperationGrant
    OperationGrant.objects.get_or_create(user=user, code=code)
    return user


@pytest.fixture
def item_instance(branch):
    from apps.categories.models import Category
    return Category.objects.create(
        asset_category='固定', item_category='办公',
        asset_name='ThinkPad T14', asset_code='NB-001', unit='台',
        management_type='instance',
    )


@pytest.fixture
def inst(branch, item_instance):
    from apps.assets.models import FixedAsset
    return FixedAsset.objects.create(
        item=item_instance, 内部编号='NB-001-1',
        当前状态='在库', branch=branch, 入库日期=datetime.date(2026, 8, 24),
    )


def _purchase_doc(branch, item, qty=1):
    """造一张已生效采购单 + 明细行（实例出生的合规路径）。"""
    from apps.transfers.models import Transfer, TransferLine
    transfer = Transfer.objects.create(
        单据编号='CG20260824-001', 调拨日期=datetime.date(2026, 8, 24),
        调出分公司=branch.name, from_branch=branch, to_branch=branch,
        action_type='purchase', 审批状态='已入库', 供应商='联想',
    )
    return TransferLine.objects.create(
        transfer=transfer, item=item, 行号=1, 数量=qty, 单价=7999,
    )


@pytest.mark.django_db
class TestFrozenWriteEndpoints:
    def test_create_frozen(self, supervisor_user):
        client = _client_for(supervisor_user)
        resp = client.post('/api/assets/fixed-assets', {'内部编号': 'X-1'})
        assert resp.status_code == 405
        assert '流转单' in resp.data['detail']

    def test_update_frozen(self, supervisor_user, inst):
        client = _client_for(supervisor_user)
        resp = client.patch(f'/api/assets/fixed-assets/{inst.pk}', {'使用人': '张三'})
        assert resp.status_code == 405

    def test_destroy_frozen(self, supervisor_user, inst):
        client = _client_for(supervisor_user)
        resp = client.delete(f'/api/assets/fixed-assets/{inst.pk}')
        assert resp.status_code == 405

    def test_batch_delete_frozen(self, supervisor_user, inst):
        client = _client_for(supervisor_user)
        resp = client.post('/api/assets/fixed-assets/batch-delete', {'ids': [str(inst.pk)]})
        assert resp.status_code == 405
        from apps.assets.models import FixedAsset
        assert FixedAsset.objects.filter(pk=inst.pk).exists()

    def test_import_gone(self, supervisor_user):
        client = _client_for(supervisor_user)
        resp = client.post('/api/assets/fixed-assets/import')
        assert resp.status_code == 410


@pytest.mark.django_db
class TestInstanceListOutput:
    def test_list_joins_dictionary_and_birth(self, supervisor_user, branch, inst):
        line = _purchase_doc(branch, inst.item)
        inst.birth_line = line
        inst.save(update_fields=['birth_line'])

        client = _client_for(supervisor_user)
        resp = client.get('/api/assets/fixed-assets')
        assert resp.status_code == 200
        data = resp.data['results'][0]
        assert data['item_code'] == 'NB-001'
        assert data['item_name'] == 'ThinkPad T14'
        assert data['management_type'] == 'instance'
        assert data['待补录'] is True          # 序列号为空
        assert data['供应商'] == '联想'         # 出生行派生
        assert str(data['单价']) == '7999.00'
        assert str(data['采购日期']) == '2026-08-24'

    def test_pending_serial_filter(self, supervisor_user, inst):
        from apps.assets.models import FixedAsset
        FixedAsset.objects.create(
            item=inst.item, 内部编号='NB-001-2', 当前状态='在库',
            branch=inst.branch, 序列号='SN-XYZ',
        )
        client = _client_for(supervisor_user)
        resp = client.get('/api/assets/fixed-assets?pending_serial=1')
        assert resp.data['count'] == 1
        assert resp.data['results'][0]['内部编号'] == 'NB-001-1'

    def test_export_headers(self, supervisor_user, inst):
        client = _client_for(supervisor_user)
        resp = client.get('/api/assets/fixed-assets/export')
        assert resp.status_code == 200
        assert FA_EXPORT_HEADERS[0] == '序号'
        assert '品目编号' in FA_EXPORT_HEADERS
        assert '电脑序列号' not in FA_EXPORT_HEADERS


@pytest.mark.django_db
class TestSupplement:
    def test_supplement_with_manage_instances(self, supervisor_user, inst):
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        resp = client.patch(
            f'/api/assets/fixed-assets/{inst.pk}/supplement',
            {'序列号': 'SN-123', '备注': '首批'},
        )
        assert resp.status_code == 200
        inst.refresh_from_db()
        assert inst.序列号 == 'SN-123'
        assert inst.备注 == '首批'
        assert resp.data['待补录'] is False

    def test_supplement_rejects_state_fields(self, supervisor_user, inst):
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        resp = client.patch(
            f'/api/assets/fixed-assets/{inst.pk}/supplement',
            {'序列号': 'SN-123', '当前状态': '在用'},
        )
        assert resp.status_code == 400

    def test_supplement_denied_without_operation(self, supervisor_user, inst):
        client = _client_for(supervisor_user)
        resp = client.patch(
            f'/api/assets/fixed-assets/{inst.pk}/supplement',
            {'序列号': 'SN-123'},
        )
        assert resp.status_code == 403


def _image_file(name='photo.jpg', content_type='image/jpeg', size_kb=1):
    from django.core.files.uploadedfile import SimpleUploadedFile
    return SimpleUploadedFile(
        name, b'\xff' * (size_kb * 1024), content_type=content_type,
    )


@pytest.mark.django_db
class TestInstanceImage:
    def test_upload_image_success(self, supervisor_user, inst):
        from django.core.files.storage import default_storage
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        resp = client.post(
            f'/api/assets/fixed-assets/{inst.pk}/image',
            {'image': _image_file()},
        )
        assert resp.status_code == 200
        assert '/media/fixed_assets/' in resp.data['图片']
        inst.refresh_from_db()
        assert inst.image.name.startswith('fixed_assets/')
        assert default_storage.exists(inst.image.name)

    def test_upload_rejects_bad_content_type(self, supervisor_user, inst):
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        resp = client.post(
            f'/api/assets/fixed-assets/{inst.pk}/image',
            {'image': _image_file('photo.gif', content_type='image/gif')},
        )
        assert resp.status_code == 400
        assert 'JPG、PNG、WebP' in resp.data['detail']
        inst.refresh_from_db()
        assert not inst.image

    def test_upload_rejects_oversize(self, supervisor_user, inst):
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        resp = client.post(
            f'/api/assets/fixed-assets/{inst.pk}/image',
            {'image': _image_file(size_kb=3 * 1024)},
        )
        assert resp.status_code == 400
        assert '2MB' in resp.data['detail']
        inst.refresh_from_db()
        assert not inst.image

    def test_upload_missing_file(self, supervisor_user, inst):
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        resp = client.post(f'/api/assets/fixed-assets/{inst.pk}/image')
        assert resp.status_code == 400

    def test_overwrite_cleans_old_file(self, supervisor_user, inst):
        from django.core.files.storage import default_storage
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        resp = client.post(
            f'/api/assets/fixed-assets/{inst.pk}/image',
            {'image': _image_file('a.jpg')},
        )
        old_name = resp.data['图片'].split('/media/')[-1]
        resp = client.post(
            f'/api/assets/fixed-assets/{inst.pk}/image',
            {'image': _image_file('b.jpg')},
        )
        assert resp.status_code == 200
        assert not default_storage.exists(old_name)
        inst.refresh_from_db()
        assert default_storage.exists(inst.image.name)

    def test_delete_image(self, supervisor_user, inst):
        from django.core.files.storage import default_storage
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        client.post(
            f'/api/assets/fixed-assets/{inst.pk}/image',
            {'image': _image_file()},
        )
        inst.refresh_from_db()
        name = inst.image.name
        resp = client.delete(f'/api/assets/fixed-assets/{inst.pk}/image')
        assert resp.status_code == 200
        assert resp.data['图片'] is None
        inst.refresh_from_db()
        assert not inst.image
        assert not default_storage.exists(name)

    def test_image_endpoints_denied_without_operation(self, supervisor_user, inst):
        client = _client_for(supervisor_user)
        resp = client.post(
            f'/api/assets/fixed-assets/{inst.pk}/image',
            {'image': _image_file()},
        )
        assert resp.status_code == 403
        resp = client.delete(f'/api/assets/fixed-assets/{inst.pk}/image')
        assert resp.status_code == 403
        inst.refresh_from_db()
        assert not inst.image

    def test_upload_ignores_state_fields(self, supervisor_user, inst):
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        resp = client.post(
            f'/api/assets/fixed-assets/{inst.pk}/image',
            {'image': _image_file(), '当前状态': '在用', '使用人': '张三'},
        )
        assert resp.status_code == 200
        inst.refresh_from_db()
        assert inst.当前状态 == '在库'
        assert inst.使用人 == ''


@pytest.mark.django_db
class TestTimeline:
    def test_timeline_covers_birth_and_flows(self, supervisor_user, branch, inst):
        from apps.transfers.models import Transfer, TransferLineInstance
        line = _purchase_doc(branch, inst.item)
        inst.birth_line = line
        inst.save(update_fields=['birth_line'])
        TransferLineInstance.objects.create(line=line, instance=inst)

        assign = Transfer.objects.create(
            单据编号='LY20260824-001', 调拨日期=datetime.date(2026, 8, 24),
            调出分公司=branch.name, from_branch=branch,
            action_type='assign', 审批状态='已通过',
        )
        assign_line = assign.lines.create(
            item=inst.item, 行号=1, 数量=1, 使用人='张三',
        )
        TransferLineInstance.objects.create(line=assign_line, instance=inst)

        client = _client_for(supervisor_user)
        resp = client.get(f'/api/assets/fixed-assets/{inst.pk}/timeline')
        assert resp.status_code == 200
        assert resp.data['birth']['供应商'] == '联想'
        assert resp.data['birth']['单据编号'] == 'CG20260824-001'
        actions = [row['action_type'] for row in resp.data['timeline']]
        assert actions == ['assign', 'purchase']  # 倒序
        assert resp.data['timeline'][0]['使用人'] == '张三'


@pytest.mark.django_db
class TestNaturalOrder:
    """实例编号自然排序（第 28 案）：跨位数 -1..-12 按数字序。"""

    def test_list_natural_order(self, supervisor_user, branch, item_instance, branch_factory=None):
        from apps.assets.models import FixedAsset
        from apps.organizations.models import Branch
        bj = Branch.objects.get_or_create(name='北京排序分公司', code='BJX01', team=branch.team)[0]
        made = []
        for i in range(1, 13):
            made.append(FixedAsset.objects.create(
                item=item_instance, 内部编号=f'NB-001-BJX01-{i}',
                当前状态='在库', branch=bj,
            ))
        client = _client_for(supervisor_user)
        resp = client.get('/api/assets/fixed-assets?branch=北京排序分公司&pageSize=20')
        codes = [r['内部编号'] for r in resp.data['results']]
        assert codes == [f'NB-001-BJX01-{i}' for i in range(1, 13)], codes

    def test_branch_item_code_order(self, supervisor_user, branch, item_instance):
        """排序：分公司 → 品目编号 → 序号自然序。"""
        from apps.assets.models import FixedAsset
        from apps.categories.models import Category
        from apps.organizations.models import Branch
        team = branch.team
        bj = Branch.objects.get_or_create(name='北京排序分公司', code='BJX01', team=team)[0]
        aa = Branch.objects.get_or_create(name='安庆排序分公司', code='AQX01', team=team)[0]
        z_item = Category.objects.create(
            asset_category='固定', item_category='办公', asset_name='Z品目',
            asset_code='ZB-999', unit='台', management_type='instance')
        for i in [2, 10, 1]:
            FixedAsset.objects.create(item=z_item, 内部编号=f'ZB-999-BJX01-{i}', 当前状态='在库', branch=bj)
        for i in [9, 1]:
            FixedAsset.objects.create(item=item_instance, 内部编号=f'NB-001-BJX01-{i}', 当前状态='在库', branch=bj)
        # 安庆（字母序在北京前）一品目一实例
        FixedAsset.objects.create(item=item_instance, 内部编号='NB-001-AQX01-1', 当前状态='在库', branch=aa)
        client = _client_for(supervisor_user)
        resp = client.get('/api/assets/fixed-assets?pageSize=20')
        codes = [r['内部编号'] for r in resp.data['results']]
        # 分公司名按 DB 默认序（UTF-8 字节序：北<安），同司内品目→序号自然序
        assert codes == [
            'NB-001-BJX01-1', 'NB-001-BJX01-9',
            'ZB-999-BJX01-1', 'ZB-999-BJX01-2', 'ZB-999-BJX01-10',
            'NB-001-AQX01-1',
        ], codes

    def test_export_natural_order(self, supervisor_user, branch, item_instance):
        import openpyxl
        from io import BytesIO
        from apps.assets.models import FixedAsset
        from apps.organizations.models import Branch
        bj = Branch.objects.get_or_create(name='北京排序分公司', code='BJX01', team=branch.team)[0]
        for i in [2, 10, 1, 11]:
            FixedAsset.objects.create(
                item=item_instance, 内部编号=f'NB-001-BJX01-{i}', 当前状态='在库', branch=bj,
            )
        client = _client_for(supervisor_user)
        resp = client.get('/api/assets/fixed-assets/export?branch=北京排序分公司')
        ws = openpyxl.load_workbook(BytesIO(resp.content)).active
        codes = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
        assert codes == ['NB-001-BJX01-1', 'NB-001-BJX01-2', 'NB-001-BJX01-10', 'NB-001-BJX01-11']


@pytest.mark.django_db
class TestItemSpecFromBirthLine:
    """实例规格：出生行本批规格优先，空回退品目字典（第 29 案）。"""

    def test_birth_spec_preferred(self, supervisor_user, branch, inst):
        line = _purchase_doc(branch, inst.item)
        line.本批规格 = 'OPPO 128G'
        line.save(update_fields=['本批规格'])
        inst.birth_line = line
        inst.save(update_fields=['birth_line'])
        client = _client_for(supervisor_user)
        resp = client.get(f'/api/assets/fixed-assets/{inst.pk}')
        assert resp.data['item_spec'] == 'OPPO 128G'

    def test_fallback_to_dictionary(self, supervisor_user, branch, inst):
        # 出生行本批规格空 → 回退字典规格
        inst.item.specification = '字典标准规格'
        inst.item.save(update_fields=['specification'])
        client = _client_for(supervisor_user)
        resp = client.get(f'/api/assets/fixed-assets/{inst.pk}')
        assert resp.data['item_spec'] == '字典标准规格'

    def test_no_birth_line_uses_dictionary(self, supervisor_user, branch, inst):
        inst.item.specification = '存量规格'
        inst.item.save(update_fields=['specification'])
        client = _client_for(supervisor_user)
        resp = client.get(f'/api/assets/fixed-assets/{inst.pk}')
        assert resp.data['item_spec'] == '存量规格'


@pytest.mark.django_db
class TestBatchUpdate:
    """实例批量维护（第 31 案）：白名单 供应商/备注/序列号。"""

    def _two(self, branch, item_instance):
        from apps.assets.models import FixedAsset
        import datetime
        from apps.transfers.models import Transfer, TransferLine
        t = Transfer.objects.create(
            单据编号='BU20260915-001', 调拨日期=datetime.date(2026, 9, 15),
            调入分公司=branch.name, to_branch=branch,
            action_type='purchase', 审批状态='已入库',
        )
        l1 = TransferLine.objects.create(transfer=t, item=item_instance, 行号=1, 数量=1)
        l2 = TransferLine.objects.create(transfer=t, item=item_instance, 行号=2, 数量=1)
        a = FixedAsset.objects.create(item=item_instance, 内部编号='BU-1', 当前状态='在库', branch=branch, birth_line=l1)
        b = FixedAsset.objects.create(item=item_instance, 内部编号='BU-2', 当前状态='在库', branch=branch, birth_line=l2)
        return a, b

    def test_batch_supplier_and_remark(self, supervisor_user, branch, item_instance):
        _grant(supervisor_user, 'manage_instances')
        a, b = self._two(branch, item_instance)
        client = _client_for(supervisor_user)
        resp = client.post('/api/assets/fixed-assets/batch-update', {
            'ids': [str(a.pk), str(b.pk)], '供应商': '联想', '备注': '首批',
        }, format='json')
        assert resp.status_code == 200 and resp.data['updated'] == 2
        # 新语义（instance-level-supplier）：供应商写实例个体字段，出生行不动
        a.refresh_from_db()
        assert a.供应商 == '联想'
        b.refresh_from_db()
        assert b.备注 == '首批'

    def test_batch_serials_pairing(self, supervisor_user, branch, item_instance):
        _grant(supervisor_user, 'manage_instances')
        a, b = self._two(branch, item_instance)
        client = _client_for(supervisor_user)
        resp = client.post('/api/assets/fixed-assets/batch-update', {
            'ids': [str(a.pk), str(b.pk)], '序列号列表': ['SN-A', ''],
        }, format='json')
        assert resp.data['updated'] == 1
        assert any('序列号为空' in e for e in resp.data['errors'])
        a.refresh_from_db()
        assert a.序列号 == 'SN-A'

    def test_rejects_unknown_fields(self, supervisor_user, branch, inst):
        _grant(supervisor_user, 'manage_instances')
        client = _client_for(supervisor_user)
        resp = client.post('/api/assets/fixed-assets/batch-update', {
            'ids': [str(inst.pk)], '当前状态': '在用',
        }, format='json')
        assert resp.status_code == 400
        assert '多余字段' in resp.data['detail']

    def test_requires_permission(self, supervisor_user, inst):
        client = _client_for(supervisor_user)
        resp = client.post('/api/assets/fixed-assets/batch-update', {
            'ids': [str(inst.pk)], '备注': 'x',
        }, format='json')
        assert resp.status_code == 403

    def test_birth_line_supplier_preferred_on_instance(self, supervisor_user, branch, item_instance):
        a, _ = self._two(branch, item_instance)
        from rest_framework.test import APIRequestFactory
        from rest_framework.request import Request
        from apps.assets.serializers import FixedAssetSerializer
        a.birth_line.供应商 = '华为'
        a.birth_line.transfer.供应商 = '单头供应商'
        a.birth_line.save(update_fields=['供应商'])
        data = FixedAssetSerializer(a).data
        assert data['供应商'] == '华为'  # 行级优先
        a.birth_line.供应商 = ''
        a.birth_line.save(update_fields=['供应商'])
        data = FixedAssetSerializer(a).data
        assert data['供应商'] == '单头供应商'  # 回退单头


@pytest.mark.django_db
class TestInstanceLevelSupplier:
    """实例级供应商覆盖（批量修改隔离）：勾选谁改谁，同行不受影响。"""

    def _three_on_one_line(self, branch, item_instance):
        import datetime
        from apps.assets.models import FixedAsset
        from apps.transfers.models import Transfer, TransferLine
        t = Transfer.objects.create(
            单据编号='ISOL-001', 调拨日期=datetime.date(2026, 9, 18),
            调入分公司=branch.name, to_branch=branch,
            action_type='purchase', 审批状态='已入库', 供应商='批次商',
        )
        line = TransferLine.objects.create(transfer=t, item=item_instance, 行号=1, 数量=3)
        made = []
        for i in range(1, 4):
            made.append(FixedAsset.objects.create(
                item=item_instance, 内部编号=f'ISOL-{i}', 当前状态='在库',
                branch=branch, birth_line=line,
            ))
        return made, line

    def test_batch_supplier_isolated(self, supervisor_user, branch, item_instance):
        _grant(supervisor_user, 'manage_instances')
        made, line = self._three_on_one_line(branch, item_instance)
        client = _client_for(supervisor_user)
        resp = client.post('/api/assets/fixed-assets/batch-update', {
            'ids': [str(made[0].pk)], '供应商': '个体商',
        }, format='json')
        assert resp.status_code == 200 and resp.data['updated'] == 1
        line.refresh_from_db()
        assert line.供应商 == ''  # 出生行未被触碰
        made[0].refresh_from_db()
        made[1].refresh_from_db()
        assert made[0].供应商 == '个体商'
        assert made[1].供应商 == ''  # 同行未勾选实例无个体覆盖
        # 派生：勾选台显示个体商；未勾选台回退批次口径（单头）
        from apps.assets.serializers import FixedAssetSerializer
        assert FixedAssetSerializer(made[0]).data['供应商'] == '个体商'
        assert FixedAssetSerializer(made[1]).data['供应商'] == '批次商'

    def test_three_level_derivation(self, branch, item_instance):
        made, line = self._three_on_one_line(branch, item_instance)
        from apps.assets.serializers import FixedAssetSerializer
        made[0].供应商 = '覆盖商'
        made[0].save(update_fields=['供应商'])
        line.供应商 = '行商'
        line.save(update_fields=['供应商'])
        assert FixedAssetSerializer(made[0]).data['供应商'] == '覆盖商'
        assert FixedAssetSerializer(made[1]).data['供应商'] == '行商'
        line.供应商 = ''
        line.save(update_fields=['供应商'])
        assert FixedAssetSerializer(made[1]).data['供应商'] == '批次商'
