"""回收单台账联动契约：去向二选一（入回收库 / 直接处置）+ immediate 即时生效。

对应 document-ledger-sync 能力的回收部分；领用/归还/调拨矩阵见 test_ledger_contract.py。
"""
import pytest
from conftest import _client_for
from rest_framework import status

from apps.assets.models import AssetStock, FixedAsset
from apps.assets.services import ledger


def _ensure_item(code, management_type='quantity'):
    from apps.categories.models import Category
    item, _ = Category.objects.get_or_create(
        asset_code=code,
        defaults={
            'asset_category': '测试类目', 'item_category': '测试分类',
            'asset_name': f'品目 {code}', 'unit': '个',
            'management_type': management_type,
        },
    )
    return item


def _seed_ledger(branch, code, in_use=10, stock=0, management_type='quantity'):
    """经调整单（唯一写入口）造台账底数。"""
    item = _ensure_item(code, management_type)
    if stock:
        ledger.apply_adjustment(branch, item, ledger.COLUMN_STOCK, stock, '测试造数')
    if in_use:
        ledger.apply_adjustment(branch, item, ledger.COLUMN_IN_USE, in_use, '测试造数')
    return AssetStock.objects.get(branch=branch, item=item)


def _recovery_payload(item_id, branch, code, qty=3, instances=None, **overrides):
    """单头 + items 明细行（P2 契约）：品目经字典 uuid 引用，实例引用在明细行 instances。"""
    line = {'item': item_id(code), '数量': qty}
    if instances is not None:
        line['instances'] = [str(pk) for pk in instances]
    payload = {
        '调拨日期': '2026-08-23',
        '调出分公司': branch.name,
        'items': [line],
    }
    payload.update(overrides)
    return payload


def _approve(client, pk):
    return client.post(f'/api/transfers/{pk}/approve', {'approved': True}, format='json')


def _row(branch, code):
    return AssetStock.objects.get(branch=branch, item__asset_code=code)


@pytest.mark.django_db
class TestRecoveryRestockAndDisposePrecheck:
    def test_quantity_restock_rejected(self, authenticated_client, branch, item_id):
        """数量品物无重新入库概念（recovery-stock-source-and-cleanup）：默认去向即拒。"""
        _seed_ledger(branch, 'RC-1', in_use=10)
        resp = authenticated_client.post(
            '/api/transfers/recovery', _recovery_payload(item_id, branch, 'RC-1', qty=3), format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '数量品物无重新入库概念' in str(resp.data)
        row = _row(branch, 'RC-1')
        assert row.在用数量 == 10 and row.总量 == 10  # 台账零变化

    def test_dispose_insufficient_stock_rejected_at_create(self, authenticated_client, branch, item_id):
        """创建端软预检：数量品处置扣在库，超在库提交即拒，不落库。"""
        _seed_ledger(branch, 'RC-2', in_use=50, stock=2)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'RC-2', qty=5, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '数量品物处置扣在库' in str(resp.data['detail'])
        assert '当前在库 2' in str(resp.data['detail'])
        assert 'RC-2' in str(resp.data['detail'])
        row = _row(branch, 'RC-2')
        assert row.在库数量 == 2 and row.在用数量 == 50  # 台账零变化

    def test_dispose_insufficient_stock_caught_by_approve_final_check(
        self, authenticated_client, branch, item_id,
    ):
        """创建时在库足够、审批前账面被扣 → 行锁终检兜底（软预检不覆盖竞态）。"""
        _seed_ledger(branch, 'RC-2F', in_use=0, stock=2)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'RC-2F', qty=2, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == 201
        # 创建后、审批前，在库被另一路径（调整单）扣光
        ledger.apply_adjustment(
            branch, _ensure_item('RC-2F'), ledger.COLUMN_STOCK, -2, '模拟账面变动',
        )
        resp = _approve(authenticated_client, resp.data['id'])
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '回收扣减不足' in str(resp.data['detail'])
        row = _row(branch, 'RC-2F')
        assert row.在库数量 == 0 and row.回收库数量 == 0  # 整体回滚

    def test_return_insufficient_keeps_generic_error(self, authenticated_client, branch, item_id):
        """归还的台账不足保持通用格式（业务化文案仅回收单）。"""
        _ensure_item('RC-5')  # 无台账行 = 在用 0
        payload = {
            '调拨日期': '2026-08-23', '调入分公司': branch.name,
            'items': [{'item': item_id('RC-5'), '数量': 1}],
        }
        resp = authenticated_client.post('/api/transfers/return', payload, format='json')
        assert resp.status_code == 201
        resp = _approve(authenticated_client, resp.data['id'])
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '在用数量不足' in str(resp.data['detail'])
        assert '需变动 -1' in str(resp.data['detail'])


@pytest.mark.django_db
class TestRecoveryDirectDispose:
    def test_dispose_drops_total_without_recycle(self, authenticated_client, branch, item_id):
        """数量品处置扣在库（在用不动），总量随处置下跌。"""
        _seed_ledger(branch, 'RC-3', in_use=6, stock=4)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'RC-3', qty=2, 回收去向='dispose', 处置方式='出售', 处置金额=500),
            format='json',
        )
        assert resp.status_code == 201
        tid = resp.data['id']
        assert resp.data['回收去向'] == 'dispose'
        resp = _approve(authenticated_client, tid)
        assert resp.status_code == 200
        row = _row(branch, 'RC-3')
        assert row.在库数量 == 2  # 处置扣在库
        assert row.在用数量 == 6  # 在用不动
        assert row.回收库数量 == 0
        assert row.总量 == 8  # 总量随处置下跌

    def test_dispose_requires_no_amount_but_records_when_sold(self, authenticated_client, branch, item_id):
        _seed_ledger(branch, 'RC-4', in_use=0, stock=3)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'RC-4', qty=1, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == 201
        assert resp.data['处置方式'] == '报废'

    def test_dispose_without_method_rejected(self, authenticated_client, branch, item_id):
        """后端处置必填（建单收口）：dispose 缺处置方式即 400。"""
        _seed_ledger(branch, 'RC-4B', in_use=0, stock=3)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'RC-4B', qty=1, 回收去向='dispose'),
            format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '处置方式' in str(resp.data['detail'])

    def test_sold_without_amount_rejected(self, authenticated_client, branch, item_id):
        """后端处置必填：出售缺金额即 400。"""
        _seed_ledger(branch, 'RC-4C', in_use=0, stock=3)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'RC-4C', qty=1, 回收去向='dispose', 处置方式='出售'),
            format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '处置金额' in str(resp.data['detail'])


@pytest.mark.django_db
class TestRecoveryFixedAsset:
    def test_fa_retired_via_instance_ref_on_approve(self, authenticated_client, branch, item_id):
        """回收绑实例 → 实例转回收库（默认去向），档案保留不删除（P2 第二刀）。"""
        _seed_ledger(branch, 'FAI-1', in_use=2, management_type='instance')
        item = _ensure_item('FAI-1', 'instance')
        inst = FixedAsset.objects.create(
            item=item, 内部编号='FAI-1-1', 当前状态='在用', branch=branch, 使用人='张三',
        )
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'FAI-1', qty=1, instances=[inst.pk]),
            format='json',
        )
        assert resp.status_code == 201
        assert _approve(authenticated_client, resp.data['id']).status_code == 200
        inst.refresh_from_db()
        assert inst.当前状态 == '在库'  # 默认去向=重新入库（回收库退役）
        assert inst.使用人 == ''
        assert FixedAsset.objects.filter(pk=inst.pk).exists()  # 档案保留

    def test_in_stock_instance_dispose_retires(self, authenticated_client, branch, item_id):
        """在库实例直接处置（recovery-stock-source-and-cleanup）：台账扣在库、实例退役。"""
        _seed_ledger(branch, 'FAI-2', in_use=0, stock=1, management_type='instance')
        item = _ensure_item('FAI-2', 'instance')
        inst = FixedAsset.objects.create(
            item=item, 内部编号='FAI-2-1', 当前状态='在库', branch=branch,
        )
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'FAI-2', qty=1, instances=[inst.pk],
                              回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == 201
        assert _approve(authenticated_client, resp.data['id']).status_code == 200
        inst.refresh_from_db()
        assert inst.当前状态 == '退役'
        row = _row(branch, 'FAI-2')
        assert row.在库数量 == 0 and row.总量 == 0

    def test_in_stock_instance_restock_rejected(self, authenticated_client, branch, item_id):
        """在库实例+重新入库：预检拒绝（在库实例无需回库）。"""
        _seed_ledger(branch, 'FAI-3', in_use=0, stock=1, management_type='instance')
        item = _ensure_item('FAI-3', 'instance')
        inst = FixedAsset.objects.create(
            item=item, 内部编号='FAI-3-1', 当前状态='在库', branch=branch,
        )
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'FAI-3', qty=1, instances=[inst.pk]),
            format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '在库实例无需回库' in str(resp.data)

    def test_fa_untouched_without_instance_ref(self, authenticated_client, branch, item_id):
        """数量管理品目处置不携带实例，实例档案不受影响。"""
        _seed_ledger(branch, 'RC-6', in_use=0, stock=2)
        item = _ensure_item('RC-6')
        inst = FixedAsset.objects.create(
            item=item, 内部编号='RC-6-1', 当前状态='在库', branch=branch,
        )
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'RC-6', qty=1, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == 201
        assert _approve(authenticated_client, resp.data['id']).status_code == 200
        inst.refresh_from_db()
        assert inst.当前状态 == '在库'


@pytest.mark.django_db
class TestImmediateChannelRemoved:
    """行内即时回收通道下线（修订 5.1）：immediate 一律 400 引导走审批流。"""

    def test_immediate_rejected_for_regular_user(self, staff_user, branch, item_id):
        _seed_ledger(branch, 'IM-1', in_use=5)
        client = _client_for(staff_user)
        resp = client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'IM-1', qty=2, immediate=True),
            format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '已下线' in str(resp.data['detail'])
        assert '审批流' in str(resp.data['detail'])
        from apps.transfers.models import Transfer
        assert Transfer.objects.filter(action_type='recovery').count() == 0  # 不落库
        row = _row(branch, 'IM-1')
        assert row.在用数量 == 5  # 台账不动

    def test_immediate_rejected_even_with_manage_assets(self, supervisor_user, branch, item_id):
        """持 manage_assets 也不豁免——通道彻底下线，不留 admin 特例。"""
        _seed_ledger(branch, 'IM-2', in_use=5)
        client = _client_for(supervisor_user)
        resp = client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'IM-2', qty=2, immediate=True),
            format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        from apps.transfers.models import Transfer
        assert Transfer.objects.filter(action_type='recovery').count() == 0

    def test_recovery_create_blocked_by_inventory_lock(self, supervisor_user, branch, item_id, db):
        """普通回收单创建同样受盘点锁（immediate 时代的锁语义在创建路径延续）。"""
        from apps.inventories.models import InventoryTask
        _seed_ledger(branch, 'IM-5', in_use=0, stock=5)
        InventoryTask.objects.create(
            name='锁库盘点', branch=branch, status='in_progress',
        )
        client = _client_for(supervisor_user)
        resp = client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'IM-5', qty=1, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '盘点' in str(resp.data['detail'])

    def test_plain_recovery_stays_pending(self, staff_user, branch, item_id):
        _seed_ledger(branch, 'IM-6', in_use=0, stock=5)
        client = _client_for(staff_user)
        resp = client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'IM-6', qty=1, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == 201
        assert resp.data['审批状态'] == '待审批'
        row = _row(branch, 'IM-6')
        assert row.在库数量 == 5  # 未审批不动账


@pytest.mark.django_db
class TestRecoveryCreateStockPrecheck:
    """回收处置创建/编辑在库软预检（recovery-stock-source-and-cleanup：数量品扣在库）。"""

    def _dispose_payload(self, item_id, branch, code, items):
        return {
            '调拨日期': '2026-08-27', '调出分公司': branch.name,
            '回收去向': 'dispose', '处置方式': '报废',
            'items': items,
        }

    def test_multi_line_same_item_merged(self, authenticated_client, branch, item_id):
        """多行同品目合并计量：单行都够、合计超出即拒。"""
        _seed_ledger(branch, 'PC-1', in_use=0, stock=3)
        payload = self._dispose_payload(item_id, branch, 'PC-1', [
            {'item': item_id('PC-1'), '数量': 2},
            {'item': item_id('PC-1'), '数量': 2},
        ])
        resp = authenticated_client.post('/api/transfers/recovery', payload, format='json')
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '当前在库 3' in str(resp.data['detail']) and '需处置 4' in str(resp.data['detail'])
        from apps.transfers.models import Transfer
        assert Transfer.objects.filter(action_type='recovery').count() == 0

    def test_missing_stock_row_counts_as_zero(self, authenticated_client, branch, item_id):
        """台账行不存在（该分公司×品目从未入账）视为在库 0，直接拒。"""
        _ensure_item('PC-2')
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'PC-2', qty=1, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '当前在库 0' in str(resp.data['detail'])

    def test_other_branch_stock_not_counted(self, authenticated_client, branch, second_branch, item_id):
        """在库挂在别的分公司：本分公司视为 0（调出分公司维度），照拒。"""
        _seed_ledger(second_branch, 'PC-3', in_use=0, stock=5)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'PC-3', qty=1, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
        assert '当前在库 0' in str(resp.data['detail'])

    def test_rejected_doc_edit_over_stock_blocked(self, authenticated_client, branch, item_id):
        """驳回后编辑：把数量改成超在库，保存被拒，单据保持原状。"""
        _seed_ledger(branch, 'PC-4', in_use=0, stock=3)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'PC-4', qty=3, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == 201
        doc_id = resp.data['id']
        authenticated_client.post(f'/api/transfers/{doc_id}/approve', {'approved': False, 'reason': '核对数量'}, format='json')
        from apps.transfers.models import Transfer
        assert Transfer.objects.get(pk=doc_id).审批状态 == '已驳回'

        edit = authenticated_client.patch(
            f'/api/transfers/{doc_id}',
            {'调拨日期': '2026-08-27', '调出分公司': branch.name,
             'items': [{'item': item_id('PC-4'), '数量': 5}]},
            format='json',
        )
        assert edit.status_code == status.HTTP_400_BAD_REQUEST
        assert '当前在库 3' in str(edit.data['detail'])
        assert Transfer.objects.get(pk=doc_id).lines.get().数量 == 3  # 原行未变

    def test_stock_only_item_dispose_succeeds(self, authenticated_client, branch, item_id):
        """只有库存没有在用（办公桌在库 3 在用 0）→ 处置放行（旧规则下被拒的缺口）。"""
        _seed_ledger(branch, 'PC-5', in_use=0, stock=3)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(item_id, branch, 'PC-5', qty=1, 回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == 201
        assert _approve(authenticated_client, resp.data['id']).status_code == 200
        row = _row(branch, 'PC-5')
        assert row.在库数量 == 2 and row.总量 == 2


@pytest.mark.django_db
class TestRecoveryDisposalLedger:
    """回收台账：直接处置物资明细流水（单据派生只读）。"""

    def _make_disposed(self, client, branch):
        from apps.assets.models import FixedAsset
        _seed_ledger(branch, 'DL-1', in_use=2, management_type='instance')
        _seed_ledger(branch, 'DL-2', in_use=0, stock=5)
        item = _ensure_item('DL-1', 'instance')
        a = FixedAsset.objects.create(
            item=item, 内部编号='DL-1-1', 当前状态='在用', branch=branch, 使用人='张三')
        b = FixedAsset.objects.create(
            item=item, 内部编号='DL-1-2', 当前状态='在用', branch=branch, 使用人='李四')
        # 实例盘处置单（2 台）+ 数量品目处置单（5 个）
        for code, qty, insts, extra in [
            ('DL-1', 2, [a.pk, b.pk], {'处置方式': '出售', '处置金额': 500}),
            ('DL-2', 5, None, {'处置方式': '报废'}),
        ]:
            resp = client.post(
                '/api/transfers/recovery',
                _recovery_payload(lambda c: _ensure_item(c).id, branch, code, qty=qty, instances=insts, **{
                    '回收去向': 'dispose', **extra}),
                format='json',
            )
            assert resp.status_code == 201, resp.data
            assert _approve(client, resp.data['id']).status_code == 200

    def test_ledger_expands_instances_and_qty_rows(self, authenticated_client, branch):
        self._make_disposed(authenticated_client, branch)
        resp = authenticated_client.get('/api/transfers/recovery-ledger')
        assert resp.status_code == 200
        rows = resp.data['results']
        assert resp.data['count'] == 3
        inst_rows = [r for r in rows if r['品目编号'] == 'DL-1']
        assert sorted(r['内部编号'] for r in inst_rows) == ['DL-1-1', 'DL-1-2']
        assert all(r['数量'] == 1 and r['处置方式'] == '出售' and r['处置金额'] == 500 for r in inst_rows)
        qty_row = next(r for r in rows if r['品目编号'] == 'DL-2')
        assert qty_row['数量'] == 5 and qty_row['内部编号'] == '' and qty_row['处置方式'] == '报废'
        assert str(resp.data['disposal_income']) == '500'  # 出售金额加总（合计卡数据源）

    def test_ledger_excludes_pending_and_filters(self, authenticated_client, branch):
        from apps.assets.models import FixedAsset
        _seed_ledger(branch, 'DL-3', in_use=0, stock=1)
        resp = authenticated_client.post(
            '/api/transfers/recovery',
            _recovery_payload(lambda c: _ensure_item(c).id, branch, 'DL-3', qty=1,
                              回收去向='dispose', 处置方式='报废'),
            format='json',
        )
        assert resp.status_code == 201  # 未审批 → 不入台账
        data = authenticated_client.get('/api/transfers/recovery-ledger').data
        assert all(r['品目编号'] != 'DL-3' for r in data['results'])

        self._make_disposed(authenticated_client, branch)
        data = authenticated_client.get(
            '/api/transfers/recovery-ledger?assetCode=DL-2').data
        assert data['count'] == 1 and data['results'][0]['品目编号'] == 'DL-2'

    def test_ledger_export(self, authenticated_client, branch):
        import openpyxl
        from io import BytesIO
        self._make_disposed(authenticated_client, branch)
        resp = authenticated_client.get('/api/transfers/recovery-ledger/export')
        assert resp.status_code == 200
        ws = openpyxl.load_workbook(BytesIO(resp.content)).active
        head = [c.value for c in ws[1]]
        assert head[0] == '日期' and '内部编号' in head
        assert ws.max_row == 6  # 表头 + 3 行流水 + 空行 + 处置收入合计行
        last = [c.value for c in ws[ws.max_row]]
        assert '处置收入合计' in str(last[0]) and last[9] == 500


@pytest.mark.django_db
class TestNormalizeRecycleBin:
    """回收库退役归一命令：预览零落库、确认搬家留痕、幂等。"""

    def _seed_recycle(self, branch):
        from apps.assets.models import FixedAsset
        _seed_ledger(branch, 'NB-1', in_use=3, management_type='instance')
        item = _ensure_item('NB-1', 'instance')
        for i in range(3):
            FixedAsset.objects.create(
                item=item, 内部编号=f'NB-1-{i + 1}', 当前状态='回收库', branch=branch)
        # 台账回收库列经调整单抬高（模拟历史入回收库）
        ledger.apply_adjustment(branch, item, ledger.COLUMN_RECYCLE, 3, '造数')

    def test_preview_then_confirm_idempotent(self, branch, capsys):
        from django.core.management import call_command
        from apps.assets.models import FixedAsset, AssetStock, LedgerAdjustment
        self._seed_recycle(branch)

        call_command('normalize_recycle_bin')  # 预览
        assert FixedAsset.objects.filter(当前状态='回收库').count() == 3
        row = AssetStock.objects.get(branch=branch, item__asset_code='NB-1')
        assert row.回收库数量 == 3

        call_command('normalize_recycle_bin', '--confirm')
        assert FixedAsset.objects.filter(当前状态='回收库').count() == 0
        assert FixedAsset.objects.filter(当前状态='在库').count() == 3
        row.refresh_from_db()
        assert (row.回收库数量, row.在库数量) == (0, 3)
        assert LedgerAdjustment.objects.filter(事由__contains='回收库退役归一').count() == 2

        call_command('normalize_recycle_bin')  # 幂等
        out = capsys.readouterr().out
        assert '无需归一' in out
