import calendar
import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO

from django.db import transaction
from django.db.models import Count, Q
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from apps.assets.models import FixedAsset
from apps.organizations.models import Branch
from apps.reconciliation.models import (
    ReconciliationComparison,
    ReconciliationIssue,
    ReconciliationIssueHistory,
    ReconciliationOrderRow,
    ReconciliationTask,
    ReconciliationUploadVersion,
)
from apps.suppliers.models import Supplier


ITEM_CODES = ('A-a00001', 'A-a00008', 'A-a00011', 'A-a00012')
HEADERS = ('订单号', '序列号', '分公司', '起租日期', '记账日期', '记账天数', '月租金', '本月金额')
CENT = Decimal('0.01')


def normalize_serial(value):
    return str(value or '').strip().upper()


def _branch_key(value):
    return re.sub(r'\s+', '', str(value or ''))


def _parse_date(value, row_number, field, errors):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or '').strip()
    if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', text):
        errors.append({'row': row_number, 'field': field, 'value': text, 'message': '日期格式必须为 YYYY-MM-DD'})
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        errors.append({'row': row_number, 'field': field, 'value': text, 'message': '日期无效'})
        return None


def _parse_money(value, row_number, field, errors, required=True):
    if value in (None, ''):
        if required:
            errors.append({'row': row_number, 'field': field, 'value': '', 'message': '金额不能为空'})
        return None
    try:
        amount = Decimal(str(value).strip())
        if not amount.is_finite() or amount < 0 or amount.as_tuple().exponent < -2:
            raise InvalidOperation
        return amount.quantize(CENT)
    except (InvalidOperation, ValueError):
        errors.append({'row': row_number, 'field': field, 'value': str(value), 'message': '金额须为非负且精确到分的数字'})
        return None


def billing_for_period(period, accounting_date, monthly_rent):
    days_in_month = calendar.monthrange(period.year, period.month)[1]
    month_start = period.replace(day=1)
    month_end = period.replace(day=days_in_month)
    if accounting_date > month_end:
        billing_days = 0
    else:
        billing_days = (month_end - max(accounting_date, month_start)).days + 1
    amount = (monthly_rent * Decimal(billing_days) / Decimal(days_in_month)).quantize(
        CENT, rounding=ROUND_HALF_UP,
    )
    return billing_days, amount


def parse_workbook(upload, period):
    """Parse and validate the entire agreed workbook before any task writes."""
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError({'detail': '无法读取 Excel 文件，请确认上传的是有效的 .xlsx 文件'}) from exc

    sheet = workbook.active
    header_values = tuple(str(value or '').strip() for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
    if header_values[:len(HEADERS)] != HEADERS:
        raise ValidationError({
            'detail': '订单模板列名或顺序不正确',
            'expectedHeaders': list(HEADERS),
            'actualHeaders': list(header_values),
        })

    branches = list(Branch.objects.all().only('id', 'name'))
    branch_map = {}
    for branch in branches:
        branch_map.setdefault(_branch_key(branch.name), []).append(branch)

    errors = []
    parsed = []
    seen_serials = {}
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, max_col=len(HEADERS), values_only=True), start=2):
        if not any(value not in (None, '') for value in values):
            continue
        values = tuple(values) + (None,) * (len(HEADERS) - len(values))
        order_no, serial, branch_text, start_raw, accounting_raw, days_raw, rent_raw, amount_raw = values[:8]
        serial_original = str(serial or '').strip()
        serial_normalized = normalize_serial(serial_original)
        if not serial_normalized:
            errors.append({'row': row_number, 'field': '序列号', 'value': '', 'message': '序列号不能为空'})
        elif serial_normalized in seen_serials:
            errors.append({'row': row_number, 'field': '序列号', 'value': serial_original,
                           'message': f'序列号重复，首次出现在第 {seen_serials[serial_normalized]} 行'})
        else:
            seen_serials[serial_normalized] = row_number

        branch_original = str(branch_text or '').strip()
        matches = branch_map.get(_branch_key(branch_original), [])
        if len(matches) != 1:
            message = '分公司名称不能为空' if not branch_original else '分公司名称未匹配或对应多个系统分公司'
            errors.append({'row': row_number, 'field': '分公司', 'value': branch_original, 'message': message})

        lease_start = _parse_date(start_raw, row_number, '起租日期', errors)
        accounting = _parse_date(accounting_raw, row_number, '记账日期', errors)
        rent = _parse_money(rent_raw, row_number, '月租金', errors)
        uploaded_days = None
        if days_raw not in (None, ''):
            try:
                uploaded_days = int(days_raw)
                if Decimal(str(days_raw)) != uploaded_days or uploaded_days < 0:
                    raise ValueError
            except (ValueError, TypeError, InvalidOperation):
                errors.append({'row': row_number, 'field': '记账天数', 'value': str(days_raw), 'message': '记账天数须为非负整数'})
                uploaded_days = None
        uploaded_amount = _parse_money(amount_raw, row_number, '本月金额', errors, required=False)
        if lease_start and accounting and lease_start > accounting:
            errors.append({'row': row_number, 'field': '记账日期', 'value': accounting.isoformat(), 'message': '记账日期不能早于起租日期'})
        if lease_start and accounting and rent is not None and len(matches) == 1 and serial_normalized:
            days, amount = billing_for_period(period, accounting, rent)
            parsed.append({
                'row_number': row_number,
                'order_number': str(order_no or '').strip(),
                'serial_original': serial_original,
                'serial_normalized': serial_normalized,
                'supplier_branch_text': branch_original,
                'branch': matches[0],
                'lease_start_date': lease_start,
                'accounting_date': accounting,
                'billing_days': days,
                'monthly_rent': rent,
                'month_amount': amount,
                'uploaded_billing_days': uploaded_days,
                'uploaded_month_amount': uploaded_amount,
            })

    workbook.close()
    if errors:
        raise ValidationError({'detail': '订单校验失败，整份文件未导入', 'errors': errors})
    if not parsed:
        raise ValidationError({'detail': '订单中没有可导入的数据行'})
    return parsed


def _prepare_reconciliation_order(supplier_id, period, upload):
    try:
        period_date = date.fromisoformat(f'{period}-01')
    except (TypeError, ValueError):
        raise ValidationError({'period': '月份格式必须为 YYYY-MM'})
    supplier = Supplier.objects.filter(pk=supplier_id).first()
    if supplier is None:
        raise ValidationError({'supplier': '供应商不存在'})
    return supplier, period_date, parse_workbook(upload, period_date)


def preview_reconciliation_order(supplier_id, period, upload):
    """Validate and summarize an order workbook without persisting reconciliation data."""
    supplier, period_date, parsed = _prepare_reconciliation_order(supplier_id, period, upload)
    branch_totals = {}
    order_keys = set()
    warnings = []
    total_amount = Decimal('0.00')
    for row in parsed:
        branch_id = str(row['branch'].pk)
        bucket = branch_totals.setdefault(branch_id, {
            'branch_id': branch_id, 'branch_name': row['branch'].name,
            'row_count': 0, 'month_amount': Decimal('0.00'),
        })
        bucket['row_count'] += 1
        bucket['month_amount'] += row['month_amount']
        order_keys.add((branch_id, row['order_number']))
        total_amount += row['month_amount']
        if (row['uploaded_billing_days'] is not None and row['uploaded_billing_days'] != row['billing_days']) or (
            row['uploaded_month_amount'] is not None and row['uploaded_month_amount'] != row['month_amount']
        ):
            warnings.append({
                'row': row['row_number'],
                'uploaded_billing_days': row['uploaded_billing_days'],
                'calculated_billing_days': row['billing_days'],
                'uploaded_month_amount': str(row['uploaded_month_amount']) if row['uploaded_month_amount'] is not None else None,
                'calculated_month_amount': str(row['month_amount']),
            })
    return {
        'supplier_id': str(supplier.pk), 'supplier_name': supplier.name,
        'period': period_date.strftime('%Y-%m'),
        'branch_count': len(branch_totals), 'order_count': len(order_keys),
        'row_count': len(parsed),
        'total_amount': str(total_amount.quantize(CENT)),
        'branch_totals': [
            {**bucket, 'month_amount': str(bucket['month_amount'].quantize(CENT))}
            for bucket in branch_totals.values()
        ],
        'calculation_warnings': warnings,
    }


def _issue_key(supplier_id, issue_type, serial, discriminator=''):
    raw = f'{supplier_id}|{issue_type}|{serial}|{discriminator}'
    return hashlib.sha256(raw.encode('utf-8')).hexdigest()


def _issue_for(task, issue_type, serial, discriminator=''):
    issue_key = _issue_key(task.supplier_id, issue_type, serial, discriminator)
    issue, created = ReconciliationIssue.objects.get_or_create(
        issue_key=issue_key,
        defaults={
            'supplier': task.supplier, 'issue_type': issue_type,
            'serial_normalized': serial, 'first_period': task.period,
            'last_period': task.period,
        },
    )
    if not created:
        old_status = issue.status
        issue.last_period = max(issue.last_period, task.period)
        issue.status = ReconciliationIssue.STATUS_OPEN
        issue.save(update_fields=['last_period', 'status', 'updated_at'])
        if old_status != issue.status:
            ReconciliationIssueHistory.objects.create(
                issue=issue, user=task.created_by, old_status=old_status,
                new_status=issue.status, old_conclusion=issue.conclusion,
                new_conclusion=issue.conclusion, old_note=issue.note, new_note=issue.note,
            )
    return issue


def _snapshot_system(instance):
    return {
        'system_instance_id': instance.pk,
        'internal_code': instance.内部编号,
        'system_serial_original': instance.序列号,
        'system_branch_id': instance.branch_id,
        'system_branch_name': instance.branch.name if instance.branch_id else '',
        'system_status': instance.当前状态,
        'item_code': instance.item.asset_code,
        'item_name': instance.item.asset_name,
    }


def _create_comparison(version, result_type, order_row=None, instance=None, issue=None):
    values = {'version': version, 'result_type': result_type, 'order_row': order_row, 'issue': issue}
    if instance is not None:
        values.update(_snapshot_system(instance))
    return ReconciliationComparison.objects.create(**values)


def build_comparison(task, version):
    order_rows = list(version.order_rows.select_related('branch').all())
    order_by_serial = {row.serial_normalized: row for row in order_rows}
    instances = list(
        FixedAsset.objects.filter(
            item__asset_code__in=ITEM_CODES,
            当前状态__in=[FixedAsset.STATUS_IN_USE, FixedAsset.STATUS_IN_STOCK],
        ).select_related('item', 'branch').order_by('内部编号')
    )
    with_serial = {}
    missing_serial = []
    for instance in instances:
        serial = normalize_serial(instance.序列号)
        if not serial:
            missing_serial.append(instance)
        else:
            with_serial.setdefault(serial, []).append(instance)

    seen_issue_ids = set()
    duplicate_serials = {serial for serial, group in with_serial.items() if len(group) > 1}
    for serial in duplicate_serials:
        issue = _issue_for(task, ReconciliationComparison.RESULT_SYSTEM_DUPLICATE, serial)
        seen_issue_ids.add(issue.pk)
        order_row = order_by_serial.get(serial)
        for instance in with_serial[serial]:
            _create_comparison(version, ReconciliationComparison.RESULT_SYSTEM_DUPLICATE, order_row, instance, issue)

    for serial, order_row in order_by_serial.items():
        if serial in duplicate_serials:
            continue
        group = with_serial.get(serial, [])
        if not group:
            issue = _issue_for(task, ReconciliationComparison.RESULT_ORDER_ONLY, serial)
            seen_issue_ids.add(issue.pk)
            _create_comparison(version, ReconciliationComparison.RESULT_ORDER_ONLY, order_row, issue=issue)
            continue
        instance = group[0]
        if instance.branch_id != order_row.branch_id:
            issue = _issue_for(task, ReconciliationComparison.RESULT_BRANCH_MISMATCH, serial)
            seen_issue_ids.add(issue.pk)
            _create_comparison(version, ReconciliationComparison.RESULT_BRANCH_MISMATCH, order_row, instance, issue)
        else:
            _create_comparison(version, ReconciliationComparison.RESULT_MATCHED, order_row, instance)

    for serial, group in with_serial.items():
        if serial in duplicate_serials:
            continue
        if serial in order_by_serial:
            continue
        issue = _issue_for(task, ReconciliationComparison.RESULT_SYSTEM_ONLY, serial)
        seen_issue_ids.add(issue.pk)
        for instance in group:
            _create_comparison(version, ReconciliationComparison.RESULT_SYSTEM_ONLY, instance=instance, issue=issue)

    for instance in missing_serial:
        issue = _issue_for(
            task, ReconciliationComparison.RESULT_SYSTEM_MISSING_SERIAL, '', str(instance.pk),
        )
        seen_issue_ids.add(issue.pk)
        _create_comparison(version, ReconciliationComparison.RESULT_SYSTEM_MISSING_SERIAL, instance=instance, issue=issue)

    previous_issues = ReconciliationIssue.objects.filter(
        supplier=task.supplier,
        status__in=[ReconciliationIssue.STATUS_OPEN, ReconciliationIssue.STATUS_NOT_RECURRED],
    ).filter(
        Q(last_period__lt=task.period) | Q(occurrences__version__task=task),
    ).exclude(pk__in=seen_issue_ids)
    for issue in previous_issues:
        if issue.status != ReconciliationIssue.STATUS_NOT_RECURRED:
            old_status = issue.status
            issue.status = ReconciliationIssue.STATUS_NOT_RECURRED
            issue.save(update_fields=['status', 'updated_at'])
            ReconciliationIssueHistory.objects.create(
                issue=issue, user=task.created_by, old_status=old_status,
                new_status=issue.status, old_conclusion=issue.conclusion,
                new_conclusion=issue.conclusion, old_note=issue.note, new_note=issue.note,
            )


@transaction.atomic
def import_reconciliation_order(supplier_id, period, upload, user):
    # Parse before mutating task state so a bad replacement leaves the active version intact.
    supplier, period_date, parsed = _prepare_reconciliation_order(supplier_id, period, upload)
    upload.seek(0)
    task = ReconciliationTask.objects.select_for_update().filter(
        supplier=supplier, period=period_date,
    ).first()
    if task and task.status == ReconciliationTask.STATUS_ARCHIVED:
        raise ValidationError({'detail': '该供应商月份任务已归档，不能覆盖'})
    if task is None:
        task = ReconciliationTask.objects.create(
            supplier=supplier, period=period_date, created_by=user,
        )

    version_number = (task.versions.order_by('-version_number').values_list('version_number', flat=True).first() or 0) + 1
    version = ReconciliationUploadVersion.objects.create(
        task=task, version_number=version_number, file=upload,
        uploaded_by=user, row_count=len(parsed),
        calculation_warnings=[
            {
                'row': row['row_number'],
                'uploaded_billing_days': row['uploaded_billing_days'],
                'calculated_billing_days': row['billing_days'],
                'uploaded_month_amount': str(row['uploaded_month_amount']) if row['uploaded_month_amount'] is not None else None,
                'calculated_month_amount': str(row['month_amount']),
            }
            for row in parsed
            if (row['uploaded_billing_days'] is not None and row['uploaded_billing_days'] != row['billing_days'])
            or (row['uploaded_month_amount'] is not None and row['uploaded_month_amount'] != row['month_amount'])
        ],
    )
    rows = [ReconciliationOrderRow(version=version, **entry) for entry in parsed]
    ReconciliationOrderRow.objects.bulk_create(rows)
    version.refresh_from_db()

    # Compare against current system facts while keeping every accepted version immutable.
    build_comparison(task, version)
    task.current_version = version
    task.save(update_fields=['current_version', 'updated_at'])
    return task, version


def update_issue(issue, user, status=None, conclusion=None, note=None):
    old = {'status': issue.status, 'conclusion': issue.conclusion, 'note': issue.note}
    if status is not None:
        if status not in (ReconciliationIssue.STATUS_OPEN, ReconciliationIssue.STATUS_CLOSED):
            raise ValidationError({'status': '状态只允许待处理或结案'})
        if status == ReconciliationIssue.STATUS_CLOSED and issue.status == ReconciliationIssue.STATUS_NOT_RECURRED:
            pass
        issue.status = status
        issue.closed_by = user if status == ReconciliationIssue.STATUS_CLOSED else None
        issue.closed_at = timezone.now() if status == ReconciliationIssue.STATUS_CLOSED else None
    if conclusion is not None:
        issue.conclusion = conclusion.strip()
    if note is not None:
        issue.note = note.strip()
    issue.handled_by = user
    issue.handled_at = timezone.now()
    issue.save()
    ReconciliationIssueHistory.objects.create(
        issue=issue, user=user,
        old_status=old['status'], new_status=issue.status,
        old_conclusion=old['conclusion'], new_conclusion=issue.conclusion,
        old_note=old['note'], new_note=issue.note,
    )
    return issue


def task_summary(task):
    version = task.current_version
    if version is None:
        return {'row_count': 0, 'counts': {}, 'total_amount': '0.00', 'branch_totals': []}
    rows = list(version.order_rows.select_related('branch'))
    comparisons = list(version.comparisons.select_related('order_row__branch', 'issue'))
    counts = dict(ReconciliationComparison.objects.filter(version=version).values_list('result_type').annotate(n=Count('id')).values_list('result_type', 'n'))
    branch_totals = {}
    for row in rows:
        key = str(row.branch_id)
        bucket = branch_totals.setdefault(key, {
            'branch_id': key, 'branch_name': row.branch.name, 'row_count': 0,
            'monthly_rent': Decimal('0.00'), 'month_amount': Decimal('0.00'),
        })
        bucket['row_count'] += 1
        bucket['monthly_rent'] += row.monthly_rent
        bucket['month_amount'] += row.month_amount
    for bucket in branch_totals.values():
        bucket['monthly_rent'] = str(bucket['monthly_rent'].quantize(CENT))
        bucket['month_amount'] = str(bucket['month_amount'].quantize(CENT))
    total = sum((row.month_amount for row in rows), Decimal('0.00'))
    return {
        'row_count': len(rows), 'counts': counts, 'total_amount': str(total.quantize(CENT)),
        'branch_totals': list(branch_totals.values()), 'comparison_count': len(comparisons),
    }
