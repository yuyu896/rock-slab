import calendar
from datetime import date
from io import BytesIO

from django.db import transaction
from django.db.models import Count
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from rest_framework import status
from rest_framework.parsers import MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.audit.decorators import audit_log
from apps.audit.utils import create_audit_log
from apps.reconciliation.models import (
    ReconciliationComparison,
    ReconciliationIssue,
    ReconciliationIssueHistory,
    ReconciliationTask,
)
from apps.reconciliation.services import (
    import_reconciliation_order,
    preview_reconciliation_order,
    task_summary,
    update_issue,
)
from apps.suppliers.models import Supplier
from core.permissions import IsAdmin


class ReconciliationAdminView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]


def _task_payload(task):
    summary = task_summary(task)
    version = task.current_version
    return {
        'id': str(task.pk), 'supplier_id': str(task.supplier_id), 'supplier_name': task.supplier.name,
        'period': task.period.strftime('%Y-%m'), 'status': task.status,
        'created_at': task.created_at, 'updated_at': task.updated_at,
        'archived_at': task.archived_at, 'current_version': version.version_number if version else None,
        'uploaded_by': version.uploaded_by.name if version else None,
        'uploaded_at': version.created_at if version else None,
        **summary,
    }


def _comparison_payload(comp):
    order = comp.order_row
    issue = comp.issue
    future_billing = False
    if order:
        month_end_day = calendar.monthrange(order.version.task.period.year, order.version.task.period.month)[1]
        month_end = order.version.task.period.replace(day=month_end_day)
        future_billing = order.accounting_date > month_end
    return {
        'id': str(comp.pk), 'result_type': comp.result_type,
        'serial': order.serial_original if order else comp.system_serial_original,
        'normalized_serial': order.serial_normalized if order else issue.serial_normalized if issue else '',
        'order_number': order.order_number if order else '',
        'supplier_branch': order.supplier_branch_text if order else '',
        'supplier_branch_id': str(order.branch_id) if order else None,
        'system_branch': comp.system_branch_name,
        'system_branch_id': str(comp.system_branch_id) if comp.system_branch_id else None,
        'internal_code': comp.internal_code,
        'system_status': comp.system_status,
        'item_code': comp.item_code,
        'item_name': comp.item_name,
        'lease_start_date': order.lease_start_date.isoformat() if order else None,
        'accounting_date': order.accounting_date.isoformat() if order else None,
        'billing_days': order.billing_days if order else None,
        'monthly_rent': str(order.monthly_rent) if order else None,
        'month_amount': str(order.month_amount) if order else None,
        'uploaded_billing_days': order.uploaded_billing_days if order else None,
        'uploaded_month_amount': str(order.uploaded_month_amount) if order and order.uploaded_month_amount is not None else None,
        'billing_status': 'not_started' if future_billing else 'billable' if order else None,
        'issue': {
            'id': str(issue.pk), 'status': issue.status, 'conclusion': issue.conclusion,
            'note': issue.note, 'first_period': issue.first_period.strftime('%Y-%m'),
            'last_period': issue.last_period.strftime('%Y-%m'),
            'handled_by': issue.handled_by.name if issue.handled_by else None,
            'handled_at': issue.handled_at,
        } if issue else None,
    }


class ReconciliationTaskListView(ReconciliationAdminView):
    def get(self, request):
        tasks = ReconciliationTask.objects.select_related('supplier', 'current_version__uploaded_by').all()
        supplier = request.query_params.get('supplier')
        period = request.query_params.get('period')
        if supplier:
            tasks = tasks.filter(supplier_id=supplier)
        if period:
            try:
                tasks = tasks.filter(period=date.fromisoformat(f'{period}-01'))
            except ValueError:
                return Response({'detail': '月份格式必须为 YYYY-MM'}, status=400)
        unresolved_groups = list(ReconciliationIssue.objects.filter(
            status__in=[ReconciliationIssue.STATUS_OPEN, ReconciliationIssue.STATUS_NOT_RECURRED],
        ).values(
            'supplier_id', 'supplier__name', 'first_period',
        ).annotate(issue_count=Count('pk', distinct=True)).order_by(
            'supplier__name', '-first_period',
        ))
        return Response({
            'suppliers': [{'id': str(s.pk), 'name': s.name} for s in Supplier.objects.order_by('name')],
            'results': [_task_payload(task) for task in tasks],
            'unresolved_summary': {
                'total_count': sum(group['issue_count'] for group in unresolved_groups),
                'groups': [{
                    'supplier_id': str(group['supplier_id']),
                    'supplier_name': group['supplier__name'],
                    'first_period': group['first_period'].strftime('%Y-%m'),
                    'count': group['issue_count'],
                } for group in unresolved_groups],
            },
        })


class ReconciliationTaskPreviewView(ReconciliationAdminView):
    parser_classes = [MultiPartParser]

    def post(self, request):
        upload = request.FILES.get('file')
        if upload is None:
            return Response({'detail': '请选择订单文件'}, status=400)
        if not upload.name.lower().endswith('.xlsx'):
            return Response({'detail': '目前仅支持 .xlsx 文件'}, status=400)
        return Response(preview_reconciliation_order(
            request.data.get('supplier'), request.data.get('period'), upload,
        ))


class ReconciliationTaskUploadView(ReconciliationAdminView):
    parser_classes = [MultiPartParser]

    @audit_log(action='import', resource_type='ReconciliationTask', description_template='导入/替换月度租赁对账订单')
    def post(self, request):
        upload = request.FILES.get('file')
        if upload is None:
            return Response({'detail': '请选择订单文件'}, status=400)
        if not upload.name.lower().endswith('.xlsx'):
            return Response({'detail': '目前仅支持 .xlsx 文件'}, status=400)
        task, _version = import_reconciliation_order(
            request.data.get('supplier'), request.data.get('period'), upload, request.user,
        )
        return Response(_task_payload(task), status=status.HTTP_200_OK)


class ReconciliationTaskDetailView(ReconciliationAdminView):
    def get(self, request, pk):
        task = get_object_or_404(
            ReconciliationTask.objects.select_related('supplier', 'current_version__uploaded_by'), pk=pk,
        )
        payload = _task_payload(task)
        if task.current_version:
            comps = task.current_version.comparisons.select_related(
                'order_row__branch', 'issue__handled_by',
            ).all()
            result_filter = request.query_params.get('result')
            branch_filter = request.query_params.get('branch')
            if result_filter:
                comps = comps.filter(result_type=result_filter)
            if branch_filter:
                comps = comps.filter(order_row__branch_id=branch_filter) | comps.filter(system_branch_id=branch_filter)
            payload['comparisons'] = [_comparison_payload(comp) for comp in comps]
            payload['versions'] = [{
                'id': str(version.pk), 'version_number': version.version_number,
                'uploaded_by': version.uploaded_by.name, 'uploaded_at': version.created_at,
                'row_count': version.row_count, 'file_name': version.file.name.rsplit('/', 1)[-1],
                'calculation_warnings': version.calculation_warnings,
            } for version in task.versions.select_related('uploaded_by').all()]
        payload['unresolved_issues'] = [{
            'id': str(issue.pk), 'issue_type': issue.issue_type,
            'serial': issue.serial_normalized, 'status': issue.status,
            'first_period': issue.first_period.strftime('%Y-%m'),
            'last_period': issue.last_period.strftime('%Y-%m'),
            'conclusion': issue.conclusion, 'note': issue.note,
        } for issue in ReconciliationIssue.objects.filter(
            supplier=task.supplier,
            status__in=[ReconciliationIssue.STATUS_OPEN, ReconciliationIssue.STATUS_NOT_RECURRED],
        ).order_by('first_period')]
        return Response(payload)

    @transaction.atomic
    def post(self, request, pk):
        task = get_object_or_404(ReconciliationTask.objects.select_for_update(), pk=pk)
        if not request.path.endswith('/archive'):
            return Response({'detail': '不支持的操作'}, status=405)
        if task.status == ReconciliationTask.STATUS_ARCHIVED:
            return Response({'detail': '任务已经归档'}, status=400)
        task.status = ReconciliationTask.STATUS_ARCHIVED
        task.archived_by = request.user
        task.archived_at = timezone.now()
        task.save(update_fields=['status', 'archived_by', 'archived_at', 'updated_at'])
        create_audit_log(request, 'update', 'ReconciliationTask', task.pk,
                         resource_name=f'{task.supplier.name} {task.period:%Y-%m}',
                         description='归档月度租赁对账任务')
        return Response(_task_payload(task))


class ReconciliationIssueView(ReconciliationAdminView):
    @audit_log(action='update', resource_type='ReconciliationIssue', description_template='更新对账差异处理记录')
    def patch(self, request, pk):
        issue = get_object_or_404(ReconciliationIssue, pk=pk)
        issue = update_issue(
            issue, request.user,
            status=request.data.get('status'),
            conclusion=request.data.get('conclusion'),
            note=request.data.get('note'),
        )
        return Response({
            'id': str(issue.pk), 'status': issue.status,
            'conclusion': issue.conclusion, 'note': issue.note,
            'handled_by': issue.handled_by.name if issue.handled_by else None,
            'handled_at': issue.handled_at,
        })

    def get(self, request, pk):
        issue = get_object_or_404(ReconciliationIssue, pk=pk)
        history = ReconciliationIssueHistory.objects.filter(issue=issue).select_related('user')
        return Response({
            'id': str(issue.pk), 'status': issue.status, 'conclusion': issue.conclusion,
            'note': issue.note,
            'history': [{
                'user': row.user.name, 'created_at': row.created_at,
                'old_status': row.old_status, 'new_status': row.new_status,
                'old_conclusion': row.old_conclusion, 'new_conclusion': row.new_conclusion,
                'old_note': row.old_note, 'new_note': row.new_note,
            } for row in history],
        })


class ReconciliationExportView(ReconciliationAdminView):
    @audit_log(action='export', resource_type='ReconciliationTask', description_template='导出月度租赁对账结果')
    def get(self, request):
        task = get_object_or_404(
            ReconciliationTask.objects.select_related('supplier', 'current_version'),
            pk=request.query_params.get('task'),
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '对账结果'
        sheet.append(['结果', '序列号', '订单号', '订单分公司', '系统分公司', '内部编号', '状态', '资产编号', '起租日期', '记账日期', '系统记账天数', '订单记账天数', '月租金', '系统本月金额', '订单本月金额', '处理结论', '备注'])
        if task.current_version:
            rows = task.current_version.comparisons.select_related('order_row__branch', 'issue').all()
            for row in rows:
                order = row.order_row
                sheet.append([
                    row.get_result_type_display(),
                    order.serial_original if order else row.system_serial_original,
                    order.order_number if order else '', order.supplier_branch_text if order else '',
                    row.system_branch_name, row.internal_code, row.system_status,
                    row.item_code, order.lease_start_date.isoformat() if order else '',
                    order.accounting_date.isoformat() if order else '',
                    order.billing_days if order else '', order.uploaded_billing_days if order else '',
                    str(order.monthly_rent) if order else '', str(order.month_amount) if order else '',
                    str(order.uploaded_month_amount) if order and order.uploaded_month_amount is not None else '',
                    row.issue.conclusion if row.issue else '', row.issue.note if row.issue else '',
                ])
        if request.query_params.get('type') == 'differences' and sheet.max_row > 1:
            header = [cell.value for cell in sheet[1]]
            result_column = header.index('结果') + 1
            for row_index in range(sheet.max_row, 1, -1):
                if sheet.cell(row_index, result_column).value == '匹配':
                    sheet.delete_rows(row_index)
        output = BytesIO()
        workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="reconciliation-{task.period:%Y-%m}.xlsx"'
        return response
